#!/usr/bin/env python3
"""Repeatable Office verification fixtures. Run: python3 scripts/office/evaluate.py"""

import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

HERE = Path(__file__).resolve().parent


class OfficeVerificationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def check(self, file, source=None):
        command = [sys.executable, str(HERE / "verify.py"), str(file)]
        if source:
            command += ["--source", str(source)]
        run = subprocess.run(command, capture_output=True, text=True, check=False)
        self.assertTrue(run.stdout, run.stderr)
        return run.returncode, json.loads(run.stdout)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx unavailable")
    def test_docx_creation_and_broken_package(self):
        from docx import Document
        file = self.root / "report.docx"
        doc = Document()
        doc.add_heading("Research findings", 1)
        doc.add_paragraph("Verified content.")
        doc.add_table(rows=2, cols=2)
        doc.save(file)
        code, report = self.check(file)
        self.assertEqual(code, 0)
        self.assertGreater(report["metrics"]["text_characters"], 0)
        self.assertEqual(report["metrics"]["tables"], 1)
        broken = self.root / "broken.docx"
        with ZipFile(file) as source, ZipFile(broken, "w") as output:
            for name in source.namelist():
                if name != "word/document.xml":
                    output.writestr(name, source.read(name))
        code, report = self.check(broken)
        self.assertEqual(code, 1)
        self.assertTrue(any("word/document.xml" in error for error in report["errors"]))

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "openpyxl unavailable")
    def test_xlsx_formulas_and_error_cells(self):
        from openpyxl import Workbook
        file = self.root / "model.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Inputs"
        sheet["A1"] = 2
        sheet["B1"] = "=A1*3"
        workbook.save(file)
        code, report = self.check(file)
        self.assertEqual(code, 0)
        self.assertEqual(report["metrics"]["formulas"], 1)
        sheet["C1"] = "#DIV/0!"
        sheet["C1"].data_type = "e"
        workbook.save(file)
        code, report = self.check(file)
        self.assertEqual(code, 1)
        self.assertEqual(report["metrics"]["error_cells"], 1)

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx unavailable")
    def test_pptx_creation_and_slide_loss(self):
        from pptx import Presentation
        source = self.root / "source.pptx"
        output = self.root / "output.pptx"
        deck = Presentation()
        for title in ("Question", "Answer"):
            slide = deck.slides.add_slide(deck.slide_layouts[0])
            slide.shapes.title.text = title
        deck.save(source)
        code, report = self.check(source)
        self.assertEqual(code, 0)
        self.assertEqual(report["metrics"]["slides"], 2)
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[0])
        slide.shapes.title.text = "Question"
        deck.save(output)
        code, report = self.check(output, source)
        self.assertEqual(code, 0)
        self.assertTrue(any("fewer slides" in warning for warning in report["warnings"]))


if __name__ == "__main__":
    unittest.main(verbosity=2)
